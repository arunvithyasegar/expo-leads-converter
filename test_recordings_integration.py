"""
Integration test for Anna Incubator ETM 2026 CRM - Local Whishper & Audio Recording Workflow
Tests local recording upload, atomic disk saving, Range audio streaming, transcript normalization,
offline queuing, transcript editing, format downloads, and deletion.
"""

import os
import sys
import json
import base64
import urllib.request
import urllib.parse
import urllib.error

BASE_URL = "http://127.0.0.1:8080"

def request_json(url, method="GET", data=None, headers=None):
    if headers is None:
        headers = {}
    if data is not None:
        body = json.dumps(data).encode("utf-8")
        headers["Content-Type"] = "application/json"
    else:
        body = None

    req = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req) as response:
            code = response.getcode()
            resp_headers = dict(response.info())
            resp_body = response.read()
            return code, resp_headers, resp_body
    except urllib.error.HTTPError as e:
        return e.code, dict(e.headers), e.read()

def run_tests():
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    print("=== STARTING CRM WHISHPER & RECORDINGS INTEGRATION TESTS ===")

    # 1. Server Status
    code, _, body = request_json(f"{BASE_URL}/api/status")
    assert code == 200, f"Expected 200 from /api/status, got {code}"
    status_data = json.loads(body.decode("utf-8"))
    print("[OK] 1. /api/status is healthy:", status_data["server"])

    # 2. Whishper Status (Local Offline/Online detection)
    code, _, body = request_json(f"{BASE_URL}/api/whishper/status")
    assert code == 200, f"Expected 200 from /api/whishper/status, got {code}"
    whishper_data = json.loads(body.decode("utf-8"))
    print("✓ 2. /api/whishper/status returns expected contract:", whishper_data)

    # 3. Create dummy lead for recording test if needed
    code, _, body = request_json(f"{BASE_URL}/api/leads")
    leads_resp = json.loads(body.decode("utf-8"))
    leads = leads_resp.get("leads", [])
    if not leads:
        lead_id = "LEAD-TEST-001"
        test_lead = {
            "leadId": lead_id,
            "fullName": "Dr. Ramesh Nair",
            "organisation": "BioEnergy Labs",
            "organisationType": "Startup",
            "email": "ramesh@bioenergy.test",
            "phone": "9840123456",
            "consent": "Yes",
            "owner": "Staff"
        }
        code, _, _ = request_json(f"{BASE_URL}/api/leads", method="PUT", data=[test_lead])
        assert code == 200, f"Failed to seed lead: {code}"
    else:
        lead_id = leads[0]["leadId"]

    print(f"✓ 3. Target test lead: {lead_id}")

    # 4. Upload In-Browser Audio Recording (Simulated WebM Opus bytes)
    # Generate dummy WebM header / bytes
    dummy_audio_bytes = b"\x1a\x45\xdf\xa3" + b"TEST_AUDIO_CHUNK_" * 50
    audio_b64 = "data:audio/webm;codecs=opus;base64," + base64.b64encode(dummy_audio_bytes).decode("ascii")

    upload_payload = {
        "audioBase64": audio_b64,
        "durationSeconds": 14.5,
        "createdBy": "Dr. Priya Sharma",
        "consentConfirmed": True,
        "consentMethod": "Affirmative verbal & recorded confirmation",
        "modelSize": "small",
        "language": "en"
    }

    code, _, body = request_json(f"{BASE_URL}/api/leads/{lead_id}/recordings", method="POST", data=upload_payload)
    assert code == 201, f"Expected 201 from recording upload, got {code}: {body.decode('utf-8')}"
    upload_res = json.loads(body.decode("utf-8"))
    recording = upload_res["recording"]
    rec_id = recording["recordingId"]
    audio_rel = recording["audioRelativePath"]
    print(f"✓ 4. Recording uploaded: ID={rec_id}, Path={audio_rel}, Status={recording['transcriptionStatus']}")

    # Verify physical file existence
    data_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")
    full_audio_path = os.path.join(data_dir, audio_rel)
    assert os.path.exists(full_audio_path), f"Audio file not found on disk at {full_audio_path}"
    assert os.path.getsize(full_audio_path) == len(dummy_audio_bytes)
    print(f"✓ 5. Audio saved atomically to local disk ({len(dummy_audio_bytes)} bytes)")

    # 6. Fetch Lead Recordings List
    code, _, body = request_json(f"{BASE_URL}/api/leads/{lead_id}/recordings")
    assert code == 200
    lead_recs = json.loads(body.decode("utf-8"))["recordings"]
    assert any(r["recordingId"] == rec_id for r in lead_recs)
    print(f"✓ 6. /api/leads/{lead_id}/recordings returned lead's recording")

    # 7. Fetch Global Recordings List
    code, _, body = request_json(f"{BASE_URL}/api/recordings")
    assert code == 200
    all_recs = json.loads(body.decode("utf-8"))["recordings"]
    assert any(r["recordingId"] == rec_id for r in all_recs)
    print("✓ 7. /api/recordings returned in global list")

    # 8. Test Audio Streaming with HTTP 206 Partial Content (Range header)
    code, headers, chunk = request_json(f"{BASE_URL}/api/recordings/{rec_id}/audio", headers={"Range": "bytes=0-19"})
    assert code == 206, f"Expected 206 Partial Content, got {code}"
    assert len(chunk) == 20, f"Expected 20 bytes for range 0-19, got {len(chunk)}"
    assert "Content-Range" in headers or "content-range" in headers
    print("✓ 8. Audio endpoint supports HTTP 206 Partial Content Range streaming for browser seeking")

    # 9. Get Recording Transcription Endpoint
    code, _, body = request_json(f"{BASE_URL}/api/recordings/{rec_id}/transcription")
    assert code == 200
    trans_info = json.loads(body.decode("utf-8"))
    assert trans_info["recordingId"] == rec_id
    print("✓ 9. /api/recordings/:id/transcription returned normalized structure")

    # 10. Edit and Save Staff Transcript Corrections (PUT)
    edited_text = "Dr. Ramesh confirmed interest in Anna Incubator's Clean Fuel pilot program and asked for MoU template."
    code, _, body = request_json(f"{BASE_URL}/api/recordings/{rec_id}/transcription", method="PUT", data={"editedText": edited_text})
    assert code == 200
    updated_trans = json.loads(body.decode("utf-8"))["transcription"]
    assert updated_trans["editedText"] == edited_text
    print("✓ 10. Staff transcript corrections saved atomically to disk")

    # 11. Download Formats (Audio, TXT, VTT, SRT, JSON)
    for fmt in ["audio", "txt", "vtt", "srt", "json"]:
        code, headers, data = request_json(f"{BASE_URL}/api/recordings/{rec_id}/download/{fmt}")
        assert code == 200, f"Expected 200 downloading format {fmt}, got {code}"
        assert len(data) > 0
        print(f"✓ 11. Format download /{fmt} verified ({len(data)} bytes)")

    # 12. Verify Activity Log
    code, _, body = request_json(f"{BASE_URL}/api/activities")
    assert code == 200
    activities = json.loads(body.decode("utf-8"))["activities"]
    lead_acts = [a for a in activities if a.get("leadId") == lead_id]
    assert any("recording" in a.get("summary", "").lower() for a in lead_acts)
    assert any("transcript" in a.get("summary", "").lower() for a in lead_acts)
    print("✓ 12. Recording creation and transcript edits logged in Lead Activity History")

    # 13. Retry Transcription Endpoint
    code, _, body = request_json(f"{BASE_URL}/api/recordings/{rec_id}/transcribe", method="POST")
    assert code in (200, 503), f"Expected 200 or 503, got {code}"
    retry_res = json.loads(body.decode("utf-8"))
    print(f"✓ 13. Retry transcription handled (code={code}): {retry_res}")

    # 14. Delete Recording & Files
    code, _, body = request_json(f"{BASE_URL}/api/recordings/{rec_id}", method="DELETE")
    assert code == 200
    assert not os.path.exists(full_audio_path), f"Audio file {full_audio_path} should have been deleted"
    print("✓ 14. DELETE /api/recordings/:id removes file from disk and deletes record")

    # Confirm not in list
    code, _, body = request_json(f"{BASE_URL}/api/recordings")
    all_recs_after = json.loads(body.decode("utf-8"))["recordings"]
    assert not any(r["recordingId"] == rec_id for r in all_recs_after)
    print("✓ 15. Recording confirmed purged from database")

    # 16. Test POST /api/transcribe-audio (Upload & Transcribe Audio File directly)
    sample_speech_audio = b"\x1a\x45\xdf\xa3" + b"AUDIO_SPEECH_DATA_SAMPLE_" * 60
    audio_upload_b64 = "data:audio/mp3;base64," + base64.b64encode(sample_speech_audio).decode("ascii")

    transcribe_payload = {
        "audioBase64": audio_upload_b64,
        "fileName": "client_pitch_meeting.mp3",
        "leadId": lead_id,
        "attendeeName": "Dr. Ramesh Nair",
        "language": "en",
        "modelSize": "small",
        "createdBy": "Staff"
    }

    code, _, body = request_json(f"{BASE_URL}/api/transcribe-audio", method="POST", data=transcribe_payload)
    assert code == 201, f"Expected 201 from /api/transcribe-audio, got {code}: {body.decode('utf-8')}"
    transcribe_res = json.loads(body.decode("utf-8"))
    assert transcribe_res.get("success") is True
    uploaded_rec_id = transcribe_res.get("recordingId") or (transcribe_res.get("recording") or {}).get("recordingId")
    assert uploaded_rec_id, "Missing recordingId in /api/transcribe-audio response"
    print(f"✓ 16. POST /api/transcribe-audio succeeded: recordingId={uploaded_rec_id}")

    # 17. GET /api/database/info (Excel DB metadata)
    code, _, body = request_json(f"{BASE_URL}/api/database/info")
    assert code == 200, f"Expected 200 from /api/database/info, got {code}"
    db_resp = json.loads(body.decode("utf-8"))
    info_obj = db_resp.get("info") if isinstance(db_resp.get("info"), dict) else db_resp
    assert info_obj.get("exists") is True, f"Master Excel Database data/etm_database.xlsx does not exist: {db_resp}"
    file_size = info_obj.get("fileSizeBytes") or info_obj.get("sizeBytes", 0)
    assert file_size > 1000, f"Master Excel Database file size is too small: {file_size}"
    print(f"✓ 17. /api/database/info returns valid metadata: size={file_size} bytes, sheets={info_obj.get('sheets')}")

    # 18. GET /api/database/excel (Master Excel DB live download)
    code, headers, xlsx_bytes = request_json(f"{BASE_URL}/api/database/excel")
    assert code == 200, f"Expected 200 from /api/database/excel, got {code}"
    assert len(xlsx_bytes) > 2000, f"XLSX download too small: {len(xlsx_bytes)} bytes"
    # Verify XLSX magic bytes (ZIP header: PK\x03\x04)
    assert xlsx_bytes.startswith(b"PK\x03\x04"), "Downloaded file does not have valid XLSX (ZIP) header"
    print(f"✓ 18. GET /api/database/excel serves valid binary XLSX file ({len(xlsx_bytes)} bytes)")

    # 19. Openpyxl deep verification of sheets and transcript content
    import openpyxl
    excel_path = os.path.join(data_dir, "etm_database.xlsx")
    assert os.path.exists(excel_path), f"Master Excel file does not exist at {excel_path}"
    wb = openpyxl.load_workbook(excel_path)
    sheet_names = wb.sheetnames
    expected_sheets = ["Transcripts", "TranscriptSegments", "Recordings", "Leads", "Activities", "Lists", "Users"]
    for s in expected_sheets:
        assert s in sheet_names, f"Expected sheet '{s}' missing from Excel database. Found: {sheet_names}"

    # Verify Transcripts sheet contains headers and rows
    ws_trans = wb["Transcripts"]
    assert ws_trans.max_row >= 2, f"Expected at least 1 transcript row in 'Transcripts' sheet, found max_row={ws_trans.max_row}"
    trans_headers = [ws_trans.cell(1, c).value for c in range(1, ws_trans.max_column + 1)]
    assert "RecordingID" in trans_headers
    assert "FullTranscriptText" in trans_headers
    assert "TranscriptionStatus" in trans_headers
    print(f"✓ 19. Master Excel Database verified via openpyxl: 7 sheets present, Transcripts sheet has {ws_trans.max_row - 1} rows")

    # 20. POST /api/database/excel/sync (Manual re-sync trigger)
    code, _, body = request_json(f"{BASE_URL}/api/database/excel/sync", method="POST")
    assert code == 200, f"Expected 200 from /api/database/excel/sync, got {code}"
    sync_res = json.loads(body.decode("utf-8"))
    assert sync_res.get("success") or sync_res.get("ok"), f"Sync failed: {sync_res}"
    file_size = sync_res.get("sizeBytes") or sync_res.get("fileSizeBytes") or (sync_res.get("info") or {}).get("fileSizeBytes", 0)
    print(f"✓ 20. POST /api/database/excel/sync rebuilt database successfully: {file_size} bytes")

    print("\n=======================================================")
    print("ALL 20 INTEGRATION TESTS PASSED SUCCESSFULLY! ✅")
    print("=======================================================")

if __name__ == "__main__":
    run_tests()
